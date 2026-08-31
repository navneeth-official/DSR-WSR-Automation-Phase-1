"""Trace hl_waste_below_text_in for slide 12 — print all intermediate values."""
from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

from pptx import Presentation

import update_delivery_status as uds
from app.services.ppt_format_extractor import _hl_waste_below_text_in, extract_slide
from app.services.ppt_layout_metrics import (
    CANONICAL_LINE_HEIGHT_EMU,
    EMU_PER_INCH,
    _VISUAL_LINE_HEIGHT_BUFFER,
    cell_margins,
    count_hl_paragraphs,
    count_visual_lines_in_hl,
    estimated_text_bottom_emu,
    hl_waste_below_text_in,
    inner_content_bottom_emu,
    rendered_text_bottom_emu,
)


def emu_in(value: int) -> float:
    return round(value / EMU_PER_INCH, 4)


def main() -> None:
    ppt = ROOT / "output" / "HEB_Delivery_Status.pptx"
    prs = Presentation(str(ppt))
    slide = prs.slides[11]  # slide_index 12 (1-based)
    hl = uds.get_highlights_shape(slide)

    r0 = hl.table.rows[0].height
    r1 = hl.table.rows[1].height
    r2 = hl.table.rows[2].height
    mar_t, mar_b = cell_margins(hl.table.cell(2, 0))

    hl_top = hl.top
    hl_height = hl.height
    hl_bottom_emu = hl_top + hl_height

    para_count = count_hl_paragraphs(hl)
    visual_lines = count_visual_lines_in_hl(hl)
    line_count = max(para_count, visual_lines)

    ref_r2 = r2
    ref_para_count = 15
    per_line_from_r2 = ref_r2 / max(ref_para_count - 2, 1)
    content_h_r2 = int(per_line_from_r2 * line_count * _VISUAL_LINE_HEIGHT_BUFFER)
    content_h_canonical = int(
        CANONICAL_LINE_HEIGHT_EMU * line_count * _VISUAL_LINE_HEIGHT_BUFFER
    )

    est_ref_r2_branch = hl_top + r0 + r1 + mar_t + content_h_r2
    est_emu = estimated_text_bottom_emu(hl, ref_para_count=ref_para_count, ref_r2=ref_r2)
    est_canonical = estimated_text_bottom_emu(
        hl,
        ref_para_count=ref_para_count,
        ref_r2=ref_r2,
        per_line_emu=CANONICAL_LINE_HEIGHT_EMU,
    )
    inner_emu = inner_content_bottom_emu(hl)
    rendered_emu = rendered_text_bottom_emu(hl, ref_para_count=ref_para_count, ref_r2=ref_r2)
    waste_fn = hl_waste_below_text_in(hl, ref_r2=ref_r2)
    waste_extractor = _hl_waste_below_text_in(hl)
    entry = extract_slide(slide, 12)

    print("=== CALL CHAIN ===")
    print("ppt_format_extractor._hl_waste_below_text_in(hl)")
    print("  ref_r2 = hl.table.rows[2].height")
    print("  -> ppt_layout_metrics.hl_waste_below_text_in(hl, ref_r2=ref_r2)")
    print()
    print("=== FORMULA (ppt_layout_metrics.hl_waste_below_text_in) ===")
    print("hl_bottom_emu  = hl.top + hl.height")
    print("text_bottom_emu = rendered_text_bottom_emu(hl, ref_r2=ref_r2)")
    print("waste_in = max((hl_bottom_emu - text_bottom_emu) / 914400, 0)")
    print()
    print("=== SLIDE 12 — HIGHLIGHTS SHAPE (PPTX) ===")
    print(f"hl.top (EMU)                 {hl_top}")
    print(f"hl.top (in)                  {emu_in(hl_top)}")
    print(f"hl.height (EMU)              {hl_height}")
    print(f"hl.height (in)               {emu_in(hl_height)}")
    print(f"row[0] r0 header (EMU/in)    {r0} / {emu_in(r0)}")
    print(f"row[1] r1 status (EMU/in)    {r1} / {emu_in(r1)}")
    print(f"row[2] r2 content (EMU/in)   {r2} / {emu_in(r2)}")
    print(f"content cell marT (EMU/in)   {mar_t} / {emu_in(mar_t)}")
    print(f"content cell marB (EMU/in)   {mar_b} / {emu_in(mar_b)}")
    print()
    print("=== STEP 1: HIGHLIGHTS BOX BOTTOM ===")
    print(f"hl_bottom_emu = {hl_top} + {hl_height} = {hl_bottom_emu}")
    print(f"hl_bottom_in  = {hl_bottom_emu} / {EMU_PER_INCH} = {emu_in(hl_bottom_emu)}")
    print()
    print("=== STEP 2: LAST TEXT LINE BOTTOM (rendered_text_bottom_emu) ===")
    print(f"count_hl_paragraphs          {para_count}")
    print(f"count_visual_lines_in_hl     {visual_lines}")
    print(f"line_count = max(...)        {line_count}")
    print()
    print("Branch inside estimated_text_bottom_emu:")
    print("  rendered_text_bottom_emu passes per_line_emu=CANONICAL_LINE_HEIGHT_EMU")
    print("  so ref_r2 is NOT used for per-line height in the actual path.")
    print(f"  CANONICAL_LINE_HEIGHT_EMU  {CANONICAL_LINE_HEIGHT_EMU} ({emu_in(CANONICAL_LINE_HEIGHT_EMU)} in)")
    print(f"  _VISUAL_LINE_HEIGHT_BUFFER { _VISUAL_LINE_HEIGHT_BUFFER}")
    print(f"  content_h (canonical path) = {CANONICAL_LINE_HEIGHT_EMU} * {line_count} * {_VISUAL_LINE_HEIGHT_BUFFER}")
    print(f"                             = {content_h_canonical} EMU ({emu_in(content_h_canonical)} in)")
    print()
    print("Hypothetical if ref_r2 branch were used instead:")
    print(f"  per_line = r2 / (ref_para_count-2) = {r2} / 13 = {per_line_from_r2:.4f} EMU")
    print(f"  content_h (ref_r2 path)    = {content_h_r2} EMU ({emu_in(content_h_r2)} in)")
    print(f"  est bottom (ref_r2 path)   = top+r0+r1+marT+content_h = {est_ref_r2_branch} EMU ({emu_in(est_ref_r2_branch)} in)")
    print()
    print("Actual estimated_text_bottom_emu (ref_r2 passed, canonical per_line inside rendered):")
    print(f"  = hl.top + r0 + r1 + marT + content_h")
    print(f"  = {hl_top} + {r0} + {r1} + {mar_t} + {content_h_canonical}")
    print(f"  = {est_canonical} EMU ({emu_in(est_canonical)} in)")
    print()
    print("inner_content_bottom_emu (cell inner floor):")
    print(f"  = hl.top + r0 + r1 + r2 - marB")
    print(f"  = {hl_top} + {r0} + {r1} + {r2} - {mar_b}")
    print(f"  = {inner_emu} EMU ({emu_in(inner_emu)} in)")
    print()
    print(f"rendered_text_bottom_emu = max(estimated, inner)")
    print(f"                       = max({est_canonical}, {inner_emu})")
    print(f"                       = {rendered_emu} EMU ({emu_in(rendered_emu)} in)")
    print()
    print("=== STEP 3: WASTE DISTANCE ===")
    diff_emu = hl_bottom_emu - rendered_emu
    diff_in = round(diff_emu / EMU_PER_INCH, 4)
    print(f"diff_emu = hl_bottom - text_bottom = {hl_bottom_emu} - {rendered_emu} = {diff_emu}")
    print(f"diff_in  = {diff_emu} / {EMU_PER_INCH} = {diff_in}")
    print()
    print("=== VERIFICATION ===")
    print(f"hl_waste_below_text_in(hl, ref_r2=r2)  = {waste_fn}")
    print(f"_hl_waste_below_text_in(hl) [extractor] = {waste_extractor}")
    print(f"extract_slide(...).hl_waste_below_text_in = {entry.get('hl_waste_below_text_in')}")
    print(f"Match 0.6156? {waste_fn == 0.6156}")


if __name__ == "__main__":
    main()
