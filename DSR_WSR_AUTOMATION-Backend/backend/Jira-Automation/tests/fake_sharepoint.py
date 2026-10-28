"""In-memory stand-in for ``app.services.sharepoint_storage`` plus small test helpers."""

from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from app.services import sharepoint_storage as storage

DECKS = "Shared Documents/WSR"
TEMPLATES = "Shared Documents/WSR Template"
DATA = "Shared Documents/WSR/_data"


class FakeSharePoint:
    """Same function signatures as ``sharepoint_storage``; data lives in a dict."""

    def __init__(self) -> None:
        self.files: dict[tuple[str, str], bytes] = {}
        self.fail_uploads = False
        self.configured = True
        self.upload_log: list[tuple[str, str]] = []

    # -- helpers for tests -------------------------------------------------
    def names(self, folder: str) -> list[str]:
        return sorted(name for (fld, name) in self.files if fld == folder)

    def put(self, folder: str, name: str, content: bytes) -> None:
        self.files[(folder, name)] = content

    def _item(self, folder: str, name: str) -> dict[str, Any]:
        size = len(self.files[(folder, name)])
        return {
            "name": name,
            "id": f"{folder}/{name}",
            "size": size,
            "eTag": f'"{size}"',
            "webUrl": f"https://sp.example/{folder.replace(' ', '%20')}/{name}",
            "lastModifiedDateTime": datetime.now(timezone.utc).isoformat(),
        }

    # -- storage API -------------------------------------------------------
    def storage_configured(self) -> bool:
        return self.configured

    def require_storage(self) -> None:
        if not self.configured:
            raise storage.StorageNotConfiguredError("SharePoint storage is not configured.")

    def list_folder(self, folder: str, *, use_cache: bool = True) -> list[dict[str, Any]]:
        self.require_storage()
        key = folder.strip("/")
        return [self._item(f, n) for (f, n) in sorted(self.files) if f == key]

    def get_item(self, folder: str, filename: str) -> dict[str, Any] | None:
        self.require_storage()
        key = (folder.strip("/"), filename)
        return self._item(*key) if key in self.files else None

    def upload_file(
        self, local_path: Path, folder: str, filename: str | None = None, *, mime: str | None = None
    ) -> dict[str, Any]:
        self.require_storage()
        if self.fail_uploads:
            raise storage.StorageError("SharePoint upload failed (503).")
        name = filename or Path(local_path).name
        self.files[(folder.strip("/"), name)] = Path(local_path).read_bytes()
        self.upload_log.append((folder.strip("/"), name))
        return self._item(folder.strip("/"), name)

    def download_file(self, folder: str, filename: str, destination: Path) -> Path | None:
        self.require_storage()
        key = (folder.strip("/"), filename)
        if key not in self.files:
            return None
        destination = Path(destination)
        destination.parent.mkdir(parents=True, exist_ok=True)
        destination.write_bytes(self.files[key])
        return destination

    def delete_file(self, folder: str, filename: str) -> bool:
        self.require_storage()
        return self.files.pop((folder.strip("/"), filename), None) is not None


def make_pptx(path: Path, slides: int = 2) -> bytes:
    """Write a tiny but valid .pptx and return its bytes."""
    from pptx import Presentation

    presentation = Presentation()
    for _ in range(slides):
        presentation.slides.add_slide(presentation.slide_layouts[5])
    path.parent.mkdir(parents=True, exist_ok=True)
    presentation.save(str(path))
    return path.read_bytes()
