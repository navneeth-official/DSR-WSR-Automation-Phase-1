"""Tests for template-agnostic header logo sync."""

from __future__ import annotations

from pathlib import Path

import pytest
from pptx import Presentation

from app.services.ppt_logo_sync import (
    _picture_has_blob,
    find_heb_logo_on_slide,
    sync_all_delivery_slide_logos,
    sync_header_pictures_from_reference,
)
from app.wsr_engine.slide_ops import clear_slide_shapes, copy_shapes_to_slide

SKELETON = Path(__file__).resolve().parents[1] / "fixtures" / "sustainment_skeleton.pptx"


def _logo_ok(slide) -> bool:
    logo = find_heb_logo_on_slide(slide)
    return logo is not None and _picture_has_blob(logo)


@pytest.mark.skipif(not SKELETON.is_file(), reason="skeleton fixture missing")
def test_copy_shapes_breaks_logo_until_template_sync():
    prs = Presentation(SKELETON)
    main_slide = prs.slides[2]
    layout = main_slide.slide_layout
    prs.slides.add_slide(layout)
    contd = prs.slides[-1]
    clear_slide_shapes(contd)
    copy_shapes_to_slide(main_slide, contd)

    assert _logo_ok(main_slide)
    assert not _logo_ok(contd)

    assert sync_header_pictures_from_reference(main_slide, contd) == 1
    assert _logo_ok(contd)


@pytest.mark.skipif(not SKELETON.is_file(), reason="skeleton fixture missing")
def test_sync_all_delivery_slide_logos():
    prs = Presentation(SKELETON)
    ref = prs.slides[2]
    layout = ref.slide_layout
    prs.slides.add_slide(layout)
    clone = prs.slides[-1]
    clear_slide_shapes(clone)
    copy_shapes_to_slide(ref, clone)

    synced = sync_all_delivery_slide_logos(prs, ref)
    assert synced >= 1
    assert _logo_ok(clone)
