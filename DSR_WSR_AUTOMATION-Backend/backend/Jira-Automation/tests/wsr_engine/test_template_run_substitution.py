"""Tests for template placeholder registry and run-level substitution."""

from __future__ import annotations

from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

from app.services.template_placeholders import (
    PlaceholderContext,
    find_placeholder_tokens,
    normalize_token_name,
)
from app.services.template_run_substitution import (
    _run_text,
    paragraph_run_spans,
    substitute_placeholders_in_paragraph,
    substitute_text,
)


def test_normalize_token_name_aliases() -> None:
    assert normalize_token_name(" Sprint Name ") == "sprint_name"
    assert normalize_token_name("In-review") == "in_review"


def test_find_placeholder_tokens() -> None:
    tokens = find_placeholder_tokens("{Status} - {Sprint Name} ({To Date} – {From Date})")
    assert tokens == ["Status", "Sprint Name", "To Date", "From Date"]


def test_substitute_text_respects_template_order() -> None:
    ctx = PlaceholderContext(
        sprint_name="Sprint 42",
        status="Done",
        from_date="01 Jan 2026",
        to_date="07 Jan 2026",
    )
    template = "{Status} - {Sprint Name} ({To Date} – {From Date})"
    assert substitute_text(template, ctx) == "Done - Sprint 42 (07 Jan 2026 – 01 Jan 2026)"


def test_substitute_week_report_dates() -> None:
    ctx = PlaceholderContext.from_report_period("2026-04-16", "2026-06-15")
    assert ctx.week_start_date == "16 Apr 2026"
    assert ctx.week_end_date == "15 Jun 2026"
    assert (
        substitute_text("Highlights of week ending { Week End Date }", ctx)
        == "Highlights of week ending 15 Jun 2026"
    )
    assert substitute_text("Period start { Week Start Date }", ctx) == "Period start 16 Apr 2026"
    assert normalize_token_name("Week End Date") == "week_end_date"
    assert normalize_token_name("Week Start Date") == "week_start_date"


def test_substitute_unknown_token_left_visible() -> None:
    ctx = PlaceholderContext(sprint_name="X")
    assert substitute_text("{Sprint Name} {Unknown}", ctx) == "X {Unknown}"


def test_context_bucket_count() -> None:
    section = {
        "sprint_name": "S1",
        "sprint_status": "Open",
        "sprint_dates": "1–7",
        "completed": ["a", "b"],
        "released": [],
        "inprogress": ["c"],
    }
    ctx = PlaceholderContext.from_section(section, bucket="completed")
    assert ctx.count == "2"
    ctx_all = PlaceholderContext.from_section(section)
    assert ctx_all.total == "3"


def _paragraph_xml(*run_texts: str) -> object:
    runs = "".join(
        f'<a:r xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        f'<a:rPr b="1"/><a:t>{t}</a:t></a:r>'
        for t in run_texts
    )
    return parse_xml(
        f'<a:p xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">{runs}</a:p>'
    )


def test_substitute_in_paragraph_preserves_per_run_rpr() -> None:
    p = _paragraph_xml("{Sprint Name}", ", {Status}")
    ctx = PlaceholderContext(sprint_name="Alpha", status="Open")
    assert substitute_placeholders_in_paragraph(p, ctx) is True
    full, spans = paragraph_run_spans(p)
    assert full == "Alpha, Open"
    assert spans[0].text == "Alpha"
    assert spans[1].text == ", Open"
    r_pr = spans[0].run_elem.find(qn("a:rPr"))
    assert r_pr is not None
    assert r_pr.get("b") == "1"


def test_substitute_split_run_date_keeps_name_run_font() -> None:
    """{ From Date } split across runs — value uses the token-name run's rPr."""
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    p = parse_xml(
        f'<a:p xmlns:a="{ns}">'
        f'<a:r><a:t>({{ </a:t></a:r>'
        f'<a:r><a:rPr><a:latin typeface="Bodoni MT Condensed"/></a:rPr><a:t>From Date</a:t></a:r>'
        f'<a:r><a:t> }} – {{ To Date }})</a:t></a:r>'
        f"</a:p>"
    )
    ctx = PlaceholderContext(from_date="04 Jun 2026", to_date="17 Jun 2026")
    assert substitute_placeholders_in_paragraph(p, ctx) is True
    runs = p.findall(qn("a:r"))
    assert _run_text(runs[0]) == "("
    assert _run_text(runs[1]) == "04 Jun 2026"
    lat = runs[1].find(qn("a:rPr")).find(qn("a:latin"))
    assert lat.get("typeface") == "Bodoni MT Condensed"
    assert _run_text(runs[2]) == " – 17 Jun 2026)"


def test_normalize_bucket_story_title_tokens() -> None:
    assert normalize_token_name("Completed Story Title") == "completed_story_title"
    assert normalize_token_name("In-Progress Story Title") == "in_progress_story_title"
    assert normalize_token_name("In-Review Count") == "in_review_count"


def test_resolve_bucket_specific_count_tokens() -> None:
    section = {
        "sprint_name": "S1",
        "sprint_status": "Open",
        "completed": ["a", "b"],
        "released": ["c"],
        "inprogress": [],
    }
    ctx = PlaceholderContext.from_section(section)
    assert substitute_text("completed: {Completed Count}", ctx) == "completed: 2"
    assert substitute_text("review: {In-Review Count}", ctx) == "review: 1"
    assert substitute_text("active: {In-Progress Count}", ctx) == "active: 0"


def test_resolve_bucket_specific_story_title() -> None:
    ctx = PlaceholderContext(extra={"story_title": "My story"})
    assert substitute_text("{Completed Story Title}", ctx) == "My story"
    assert substitute_text("{In-Review Story Title}", ctx) == "My story"
