"""Tests for template block spec and content filler."""

from __future__ import annotations

import copy

from pptx.oxml import parse_xml
from pptx.oxml.ns import qn

from app.services.template_block_spec import (
    ParagraphTemplate,
    TemplateBlockSpec,
    extract_block_spec_from_cell,
)
from app.services.template_content_filler import fill_highlights_cell
from app.services.template_placeholders import find_placeholder_tokens
from app.services.template_run_substitution import paragraph_plain_text


def _cell_with_paragraphs(
    paragraph_specs: list[tuple[str, int, str] | tuple[str, int, str, dict]],
) -> object:
    """Build a minimal table cell mock with paragraph XML."""
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    paras_xml = ""
    for spec in paragraph_specs:
        if len(spec) == 4:
            text, level, bullet_font, spacing = spec
        else:
            text, level, bullet_font = spec
            spacing = {}
        ppr = f'<a:pPr lvl="{level}">'
        if bullet_font:
            ppr += f'<a:buFont typeface="{bullet_font}"/>'
        if spacing.get("spc_bef"):
            ppr += f'<a:spcBef><a:spcPts val="{spacing["spc_bef"]}"/></a:spcBef>'
        if spacing.get("ln_spc"):
            ppr += f'<a:lnSpc><a:spcPts val="{spacing["ln_spc"]}"/></a:lnSpc>'
        ppr += "</a:pPr>"
        if text:
            paras_xml += (
                f'<a:p xmlns:a="{ns}">{ppr}'
                f'<a:r><a:t>{text}</a:t></a:r></a:p>'
            )
        else:
            paras_xml += f'<a:p xmlns:a="{ns}">{ppr}</a:p>'

    class FakeTextFrame:
        def __init__(self, body):
            self._txBody = body

    class FakeCell:
        def __init__(self, body):
            self.text_frame = FakeTextFrame(body)

    body = parse_xml(f'<a:txBody xmlns:a="{ns}">{paras_xml}</a:txBody>')
    return FakeCell(body)


def test_block_spec_detects_placeholders_and_static_current_week() -> None:
    cell = _cell_with_paragraphs([
        ("Sprint – {Sprint Name}, {Status} ({From Date} – {To Date})", 0, ""),
        ("This week sprint status", 0, ""),
        ("Stories completed this week– {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
    ])
    spec = extract_block_spec_from_cell(cell)
    assert spec.sprint_line is not None
    assert "Sprint Name" in spec.sprint_line.placeholders
    assert spec.current_week is not None
    assert "this week sprint status" in spec.current_week.text_preview.lower()
    assert "completed" in spec.category_headers
    assert spec.story_bullets.get("completed") is not None
    assert "Story Title" in spec.story_bullets["completed"].placeholders


def test_block_spec_user_layout_category_before_story_any_level() -> None:
    """Category header above {Story Title} at any indent — no lvl 7 required."""
    cell = _cell_with_paragraphs([
        ("Sprint – {Sprint Name}, {Status}", 0, ""),
        ("Current week sprint status", 0, ""),
        ("Stories completed this week – {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
        ("Stories released for Partner review – {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
        ("Stories in-progress this week – {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
    ])
    spec = extract_block_spec_from_cell(cell)
    assert spec.category_headers["completed"].text_preview.startswith("Stories completed")
    assert spec.category_headers["released"].text_preview.startswith("Stories released")
    assert spec.category_headers["inprogress"].text_preview.startswith("Stories in-progress")
    assert spec.story_bullets["completed"] is not None
    assert spec.story_bullets["released"] is not None
    assert spec.story_bullets["inprogress"] is not None


def test_fill_story_title_placeholder_replaced() -> None:
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    sprint_p = parse_xml(f'<a:p xmlns:a="{ns}"><a:r><a:t>{{Sprint Name}}</a:t></a:r></a:p>')
    current_week_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:r><a:t>Weekly sprint status</a:t></a:r></a:p>'
    )
    cat_p = parse_xml(
        f'<a:p xmlns:a="{ns}">'
        f'<a:r><a:t>Stories finished this week – {{Count}} stories</a:t></a:r></a:p>'
    )
    story_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:r><a:t>{{Story Title}}</a:t></a:r></a:p>'
    )

    block_spec = TemplateBlockSpec(
        sprint_line=ParagraphTemplate("sprint_line", sprint_p, frozenset({"Sprint Name"})),
        current_week=ParagraphTemplate("current_week", current_week_p),
        category_headers={
            "completed": ParagraphTemplate(
                "category_completed", cat_p, frozenset({"Count"}), bucket="completed"
            ),
        },
        story_bullet=ParagraphTemplate("story_bullet", story_p, frozenset({"Story Title"})),
    )

    cell = _cell_with_paragraphs([])
    section = {
        "sprint_name": "Alpha",
        "sprint_status": "Open",
        "completed": ["My Jira story"],
        "released": [],
        "inprogress": [],
    }
    fill_highlights_cell(cell, block_spec, [section])

    texts = [paragraph_plain_text(p) for p in cell.text_frame._txBody.findall(qn("a:p"))]
    assert texts[1] == "Weekly sprint status"
    assert "Stories finished this week" in texts[2]
    assert texts[3] == "My Jira story"
    assert "{Story Title}" not in texts[3]


def test_fill_preserves_static_and_substitutes_placeholders() -> None:
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    sprint_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:r><a:t>{{Status}} - {{Sprint Name}}</a:t></a:r></a:p>'
    )
    current_week_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:r><a:t>Current week sprint status</a:t></a:r></a:p>'
    )
    cat_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:pPr lvl="7"/>'
        f'<a:r><a:t>Stories completed this week– {{Count}} stories</a:t></a:r></a:p>'
    )
    story_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:pPr lvl="1"/>'
        f'<a:r><a:t>₋</a:t></a:r><a:r><a:t>\tSample</a:t></a:r></a:p>'
    )

    block_spec = TemplateBlockSpec(
        sprint_line=ParagraphTemplate("sprint_line", sprint_p, frozenset({"Status", "Sprint Name"})),
        current_week=ParagraphTemplate("current_week", current_week_p),
        category_headers={
            "completed": ParagraphTemplate(
                "category_completed",
                cat_p,
                frozenset({"Count"}),
                bucket="completed",
            ),
        },
        story_bullet=ParagraphTemplate("story_bullet", story_p),
    )

    cell = _cell_with_paragraphs([])
    section = {
        "sprint_name": "Alpha",
        "sprint_status": "Open",
        "sprint_dates": "1–7 Jan",
        "completed": ["Done task"],
        "released": [],
        "inprogress": [],
    }
    fill_highlights_cell(cell, block_spec, [section], track_name="Track A")

    paras = cell.text_frame._txBody.findall(qn("a:p"))
    texts = [paragraph_plain_text(p) for p in paras]
    assert texts[0] == "Open - Alpha"
    assert texts[1] == "Current week sprint status"
    assert "Stories completed this week" in texts[2]
    assert "1" in texts[2]
    assert texts[3].endswith("Done task") or "Done task" in texts[3]


def _spc_bef_pt(p_elem) -> float | None:
    ppr = p_elem.find(qn("a:pPr"))
    if ppr is None:
        return None
    spc = ppr.find(qn("a:spcBef"))
    if spc is None:
        return None
    pt = spc.find(qn("a:spcPts"))
    if pt is None:
        return None
    return int(pt.get("val", 0)) / 100


def test_block_spec_preserves_blank_spacers_in_order() -> None:
    cell = _cell_with_paragraphs([
        ("Sprint – {Sprint Name}, {Status}", 0, ""),
        ("This is it", 0, ""),
        ("", 0, "", {"spc_bef": 235}),
        ("This week sprint status", 0, ""),
        ("Stories completed this week– {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
    ])
    spec = extract_block_spec_from_cell(cell)
    roles = [p.role for p in spec.ordered_paragraphs]
    assert roles == [
        "sprint_line",
        "current_week",
        "blank",
        "static",
        "category_completed",
        "story_bullet",
    ]
    assert spec.ordered_paragraphs[2].role == "blank"


def test_fill_replays_template_order_and_spacing() -> None:
    cell = _cell_with_paragraphs([
        ("{Status}, Sprint – {Sprint Name}", 0, ""),
        ("This is it", 0, ""),
        ("", 0, "", {"spc_bef": 235}),
        ("This week sprint status", 0, ""),
        ("Stories completed this week– {Count} stories", 1, ""),
        ("{Story Title}", 2, "", {"spc_bef": 95}),
    ])
    spec = extract_block_spec_from_cell(cell)
    out = _cell_with_paragraphs([])
    section = {
        "sprint_name": "Atlas",
        "sprint_status": "Ended",
        "completed": ["Story A"],
        "released": [],
        "inprogress": [],
    }
    fill_highlights_cell(out, spec, [section])

    paras = out.text_frame._txBody.findall(qn("a:p"))
    texts = [paragraph_plain_text(p) for p in paras]
    assert texts[1] == "This is it"
    assert texts[2] == ""
    assert _spc_bef_pt(paras[2]) == 2.35
    assert texts[3] == "This week sprint status"
    assert "Story A" in texts[5]
    assert _spc_bef_pt(paras[5]) == 0.95


def test_fill_skips_blank_after_empty_completed_bucket() -> None:
    cell = _cell_with_paragraphs([
        ("{Sprint Name}", 0, ""),
        ("This week sprint status", 0, ""),
        ("Stories completed this week– {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
        ("", 0, "", {"spc_bef": 100}),
        ("Stories released for Partner review – {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
    ])
    spec = extract_block_spec_from_cell(cell)
    out = _cell_with_paragraphs([])
    section = {
        "sprint_name": "S1",
        "sprint_status": "Open",
        "completed": [],
        "released": ["Released story"],
        "inprogress": [],
    }
    fill_highlights_cell(out, spec, [section])

    paras = out.text_frame._txBody.findall(qn("a:p"))
    texts = [paragraph_plain_text(p) for p in paras]
    assert texts.count("") == 0
    assert "completed this week" not in " ".join(texts).lower()
    assert "released story" in " ".join(texts).lower()
    cw_idx = texts.index("This week sprint status")
    rel_idx = next(i for i, t in enumerate(texts) if "released for partner" in t.lower())
    assert rel_idx == cw_idx + 1


def test_fill_skips_blank_before_empty_bucket() -> None:
    cell = _cell_with_paragraphs([
        ("{Sprint Name}", 0, ""),
        ("This week sprint status", 0, ""),
        ("Stories completed this week– {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
        ("", 0, "", {"spc_bef": 100}),
        ("Stories released for Partner review – {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
        ("", 0, "", {"spc_bef": 100}),
        ("Stories in-progress this week – {Count} stories", 1, ""),
        ("{Story Title}", 2, ""),
    ])
    spec = extract_block_spec_from_cell(cell)
    out = _cell_with_paragraphs([])
    section = {
        "sprint_name": "S1",
        "sprint_status": "Open",
        "completed": [],
        "released": [],
        "inprogress": ["Only this"],
    }
    fill_highlights_cell(out, spec, [section])

    texts = [paragraph_plain_text(p) for p in out.text_frame._txBody.findall(qn("a:p"))]
    joined = " ".join(texts).lower()
    assert "completed this week" not in joined
    assert "only this" in joined


def test_block_spec_bucket_specific_tokens_any_order() -> None:
    """In-progress before completed — buckets from token names, not position."""
    cell = _cell_with_paragraphs([
        ("Sprint – {Sprint Name}, {Status}", 0, ""),
        ("This week sprint status", 0, ""),
        ("Stories in-progress this week – {In-Progress Count} stories", 1, ""),
        ("{In-Progress Story Title}", 2, ""),
        ("Stories completed this week – {Completed Count} stories", 1, ""),
        ("{Completed Story Title}", 2, ""),
        ("Stories released for Partner review – {In-Review Count} stories", 1, ""),
        ("{In-Review Story Title}", 2, ""),
    ])
    spec = extract_block_spec_from_cell(cell)
    assert spec.category_headers["inprogress"].bucket == "inprogress"
    assert spec.category_headers["completed"].bucket == "completed"
    assert spec.category_headers["released"].bucket == "released"
    assert spec.story_bullets["inprogress"].placeholders == frozenset({"In-Progress Story Title"})
    assert spec.story_bullets["completed"].placeholders == frozenset({"Completed Story Title"})
    assert spec.story_bullets["released"].placeholders == frozenset({"In-Review Story Title"})


def test_fill_bucket_specific_story_tokens_any_order() -> None:
    cell = _cell_with_paragraphs([
        ("{Sprint Name}", 0, ""),
        ("This week sprint status", 0, ""),
        ("Stories in-progress – {In-Progress Count} stories", 1, ""),
        ("{In-Progress Story Title}", 2, ""),
        ("Stories completed – {Completed Count} stories", 1, ""),
        ("{Completed Story Title}", 2, ""),
        ("Stories released – {In-Review Count} stories", 1, ""),
        ("{In-Review Story Title}", 2, ""),
    ])
    spec = extract_block_spec_from_cell(cell)
    out = _cell_with_paragraphs([])
    section = {
        "sprint_name": "S1",
        "sprint_status": "Open",
        "completed": ["Done item"],
        "released": ["Review item"],
        "inprogress": ["Active item"],
    }
    fill_highlights_cell(out, spec, [section])

    texts = [paragraph_plain_text(p) for p in out.text_frame._txBody.findall(qn("a:p"))]
    joined = " ".join(texts)
    assert "Active item" in joined
    assert "Done item" in joined
    assert "Review item" in joined
    assert "{In-Progress Story Title}" not in joined
    assert "{Completed Story Title}" not in joined
    assert "{In-Review Story Title}" not in joined
    # Counts substituted in category headers
    assert "1" in joined


def test_empty_bucket_omitted() -> None:
    ns = "http://schemas.openxmlformats.org/drawingml/2006/main"
    sprint_p = parse_xml(f'<a:p xmlns:a="{ns}"><a:r><a:t>{{Sprint Name}}</a:t></a:r></a:p>')
    current_week_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:r><a:t>Current week sprint status</a:t></a:r></a:p>'
    )
    cat_completed = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:pPr lvl="7"/>'
        f'<a:r><a:t>Stories completed this week– {{Count}} stories</a:t></a:r></a:p>'
    )
    cat_inprogress = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:pPr lvl="7"/>'
        f'<a:r><a:t>Stories in-progress – {{Count}} stories</a:t></a:r></a:p>'
    )
    story_p = parse_xml(
        f'<a:p xmlns:a="{ns}"><a:pPr lvl="1"/><a:r><a:t>Story</a:t></a:r></a:p>'
    )

    block_spec = TemplateBlockSpec(
        sprint_line=ParagraphTemplate("sprint_line", sprint_p, frozenset({"Sprint Name"})),
        current_week=ParagraphTemplate("current_week", current_week_p),
        category_headers={
            "completed": ParagraphTemplate("category_completed", cat_completed, bucket="completed"),
            "inprogress": ParagraphTemplate("category_inprogress", cat_inprogress, bucket="inprogress"),
        },
        story_bullet=ParagraphTemplate("story_bullet", story_p),
    )

    cell = _cell_with_paragraphs([])
    section = {
        "sprint_name": "S1",
        "sprint_status": "Open",
        "completed": [],
        "released": [],
        "inprogress": ["Active work"],
    }
    fill_highlights_cell(cell, block_spec, [section])

    texts = [paragraph_plain_text(p) for p in cell.text_frame._txBody.findall(qn("a:p"))]
    joined = " ".join(texts).lower()
    assert "completed this week" not in joined
    assert "in-progress" in joined or "in progress" in joined
    assert "active work" in joined.lower()
