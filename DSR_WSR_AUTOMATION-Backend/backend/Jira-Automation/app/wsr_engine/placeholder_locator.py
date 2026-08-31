"""Geometric placeholder discovery for Highlights and Key Activities."""

from __future__ import annotations

import logging
from dataclasses import dataclass

from pptx.oxml.ns import qn
from pptx.slide import Slide

from app.services.ppt_shape_utils import (
    get_highlights_shape,
    get_key_activities_shape,
    is_highlights_table,
    iter_all_shapes,
)

logger = logging.getLogger(__name__)


@dataclass
class PlaceholderTarget:
    shape: object
    row: int
    col: int
    content_top_emu: int
    content_height_emu: int
    content_width_emu: int


def _table_header_row(table, needles: tuple[str, ...]) -> int | None:
    for ri in range(len(table.rows)):
        try:
            text = table.cell(ri, 0).text.strip().lower()
        except (IndexError, AttributeError):
            continue
        if any(n in text for n in needles):
            return ri
    return None


def locate_highlights_placeholder(slide: Slide) -> PlaceholderTarget | None:
    for shape in iter_all_shapes(slide.shapes):
        if not shape.has_table:
            continue
        if not is_highlights_table(shape):
            try:
                header = shape.table.cell(0, 0).text.strip().lower()
                if "highlights" not in header and "overall status" not in header:
                    continue
            except (IndexError, AttributeError):
                continue

        table = shape.table
        header_row = _table_header_row(table, ("highlights", "overall status")) or 0
        content_row = min(header_row + 2, len(table.rows) - 1)

        row_heights = [table.rows[i].height for i in range(len(table.rows))]
        content_top = int(shape.top + sum(row_heights[:content_row]))
        content_h = row_heights[content_row] if content_row < len(row_heights) else shape.height

        return PlaceholderTarget(
            shape=shape,
            row=content_row,
            col=0,
            content_top_emu=content_top,
            content_height_emu=int(content_h),
            content_width_emu=int(shape.width),
        )

    try:
        hl = get_highlights_shape(slide)
        table = hl.table
        content_row = 2 if len(table.rows) > 2 else len(table.rows) - 1
        row_heights = [table.rows[i].height for i in range(len(table.rows))]
        content_top = int(hl.top + sum(row_heights[:content_row]))
        return PlaceholderTarget(
            shape=hl,
            row=content_row,
            col=0,
            content_top_emu=content_top,
            content_height_emu=int(row_heights[content_row]),
            content_width_emu=int(hl.width),
        )
    except ValueError:
        logger.warning("Highlights placeholder not found on slide")
        return None


def locate_ka_placeholder(slide: Slide) -> PlaceholderTarget | None:
    ka = get_key_activities_shape(slide)
    if ka is not None:
        table = ka.table
        content_row = 1 if len(table.rows) > 1 else 0
        row_heights = [table.rows[i].height for i in range(len(table.rows))]
        content_top = int(ka.top + sum(row_heights[:content_row]))
        return PlaceholderTarget(
            shape=ka,
            row=content_row,
            col=0,
            content_top_emu=content_top,
            content_height_emu=int(row_heights[content_row]),
            content_width_emu=int(ka.width),
        )
    return None


def highlights_content_cell(slide: Slide):
    target = locate_highlights_placeholder(slide)
    if target is None:
        raise ValueError("Highlights content cell not found")
    return target.shape.table.cell(target.row, target.col)


def count_template_paragraph_slots(cell) -> int:
    paras = cell.text_frame._txBody.findall(qn("a:p"))
    return max(len(paras), 1)
