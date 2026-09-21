"""Generate IT request Word document for Entra ID / SharePoint Graph access."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

REPO_ROOT = Path(__file__).resolve().parents[1]
OUTPUT = REPO_ROOT / "docs" / "ENTRA_SHAREPOINT_IT_ACCESS_REQUEST.docx"


def _add_bullet(doc: Document, text: str, bold_prefix: str | None = None) -> None:
    p = doc.add_paragraph(style="List Bullet")
    if bold_prefix:
        run = p.add_run(bold_prefix)
        run.bold = True
        p.add_run(text)
    else:
        p.add_run(text)


def main() -> None:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_heading("Entra ID & Microsoft Graph Access Request", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    sub = doc.add_paragraph("DSR-WSR Automation — SharePoint Integration")
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub.runs[0].italic = True

    doc.add_paragraph()

    doc.add_heading("Project summary", level=1)
    doc.add_paragraph(
        "DSR-WSR Automation is an internal web application that automates Weekly "
        "Status Report (WSR) creation for delivery teams. It imports Jira/Rovo story "
        "data, applies company WSR PowerPoint templates, and generates finished "
        ".pptx decks through a React frontend and FastAPI backend hosted on a GCP "
        "Ubuntu VM. The SharePoint integration will store WSR templates and generated "
        "reports in a shared document library so teams can access them centrally "
        "without manual file transfer from the server."
    )

    doc.add_paragraph()

    doc.add_paragraph(
        "Dear IT Support Team,"
    )
    doc.add_paragraph(
        "We are integrating DSR-WSR Automation with our company SharePoint document "
        "library so the application can read WSR templates and write generated WSR "
        "PowerPoint files to an approved folder. We request your support to configure "
        "Microsoft Entra ID and Microsoft Graph access for this integration."
    )

    doc.add_heading("Why Entra ID access is required", level=1)
    doc.add_paragraph(
        "Our application runs on a server (GCP VM) and must upload and download "
        ".pptx files programmatically. Microsoft does not allow direct SharePoint "
        "access using only a username and password. All automated access must go "
        "through:"
    )
    _add_bullet(doc, "An Entra ID application registration (application identity)")
    _add_bullet(doc, "Microsoft Graph API with approved permissions and admin consent")
    doc.add_paragraph(
        "Without this, the server cannot read from or write to SharePoint in a "
        "secure, supported way."
    )
    doc.add_paragraph("Business need:")
    _add_bullet(doc, "Store generated WSR reports in a shared SharePoint folder (not personal storage)")
    _add_bullet(doc, "Enable team access via standard SharePoint permissions and audit")
    _add_bullet(doc, "Support unattended operation on the server (no manual upload after each report)")

    doc.add_heading("Planned SharePoint use", level=1)
    table = doc.add_table(rows=3, cols=3)
    table.style = "Table Grid"
    headers = ("Operation", "Target (example)", "Purpose")
    for i, h in enumerate(headers):
        table.rows[0].cells[i].text = h
        for run in table.rows[0].cells[i].paragraphs[0].runs:
            run.bold = True
    table.rows[1].cells[0].text = "Write"
    table.rows[1].cells[1].text = "SharePoint site / library / Generated/"
    table.rows[1].cells[2].text = "Save generated WSR .pptx files"
    table.rows[2].cells[0].text = "Read"
    table.rows[2].cells[1].text = "SharePoint site / library / Templates/"
    table.rows[2].cells[2].text = "Load WSR template .pptx files (planned)"
    doc.add_paragraph()
    doc.add_paragraph(
        "Only file read/write in the designated folders — no access to email, "
        "calendar, Teams, or other Microsoft 365 data."
    )

    doc.add_heading("Entra ID / Graph configuration requested", level=1)
    doc.add_paragraph(
        "Please register (or approve) an Entra application for DSR-WSR-Automation "
        "with the following:"
    )

    doc.add_heading("1. Application registration", level=2)
    _add_bullet(doc, " — DSR-WSR-Automation", bold_prefix="Name")
    _add_bullet(doc, " — Single-tenant (our organization only)", bold_prefix="Type")
    _add_bullet(
        doc,
        " — Application (daemon) with client secret or certificate for server use on GCP VM",
        bold_prefix="Authentication",
    )

    doc.add_heading("2. Microsoft Graph API permissions (application permissions)", level=2)
    perm_table = doc.add_table(rows=3, cols=2)
    perm_table.style = "Table Grid"
    perm_table.rows[0].cells[0].text = "Permission"
    perm_table.rows[0].cells[1].text = "Purpose"
    for cell in perm_table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    perm_table.rows[1].cells[0].text = "Sites.Selected (preferred)"
    perm_table.rows[1].cells[1].text = "Read/write only on the SharePoint site(s) you approve"
    perm_table.rows[2].cells[0].text = "Sites.ReadWrite.All (alternative)"
    perm_table.rows[2].cells[1].text = "If Sites.Selected cannot be used — broader site access"
    doc.add_paragraph()

    doc.add_heading("3. Admin consent", level=2)
    doc.add_paragraph("Grant admin consent for the organization on the above permissions.")

    doc.add_heading("4. SharePoint site access", level=2)
    _add_bullet(
        doc,
        "Grant the application (service principal) access to the specific SharePoint site and document library.",
    )
    _add_bullet(doc, "Confirm folder paths for Templates (read) and Generated (write).")

    doc.add_heading("5. Details to provide to our team", level=2)
    details_table = doc.add_table(rows=7, cols=2)
    details_table.style = "Table Grid"
    details_table.rows[0].cells[0].text = "Item"
    details_table.rows[0].cells[1].text = "Purpose"
    for cell in details_table.rows[0].cells:
        for run in cell.paragraphs[0].runs:
            run.bold = True
    rows = [
        ("Application (client) ID", "App identity in server configuration"),
        ("Directory (tenant) ID", "Organization identity in server configuration"),
        ("Client secret (or certificate)", "Server authentication (stored securely on VM)"),
        ("SharePoint site ID", "Microsoft Graph API target"),
        ("Document library name", "e.g. Documents or WSR"),
        ("Folder paths", "e.g. Templates/, Generated/"),
    ]
    for i, (item, purpose) in enumerate(rows, start=1):
        details_table.rows[i].cells[0].text = item
        details_table.rows[i].cells[1].text = purpose

    doc.add_heading("Security notes", level=1)
    _add_bullet(doc, "Credentials will be stored only on the server (.env or secret manager), not in source code.")
    _add_bullet(doc, "We request least privilege (Sites.Selected on one site where possible).")
    _add_bullet(doc, "No user passwords will be stored in the application.")
    _add_bullet(
        doc,
        "This request is for server-to-SharePoint file access only; it is separate from any SSO/IAP/Cognito requirements for the web UI, if applicable.",
    )

    doc.add_heading("Approval requested", level=1)
    doc.add_paragraph("Please confirm approval to:")
    _add_bullet(doc, "Register the Entra application and enable the Graph permissions above")
    _add_bullet(doc, "Grant admin consent and SharePoint site access for the designated library/folders")
    _add_bullet(doc, "Provide client ID, tenant ID, client secret, and SharePoint site details via a secure channel")

    doc.add_paragraph()
    doc.add_paragraph("Application: DSR-WSR Automation")
    doc.add_paragraph("Contact: [Your name, email, team]")
    doc.add_paragraph("Target environment: GCP Ubuntu VM (production)")
    doc.add_paragraph()
    doc.add_paragraph("Thank you for your support.")
    doc.add_paragraph()
    doc.add_paragraph("Regards,")
    doc.add_paragraph("[Your name]")
    doc.add_paragraph("[Team / project]")

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    main()
